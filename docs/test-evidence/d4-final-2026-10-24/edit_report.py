from pathlib import Path
from copy import deepcopy
import json, re, shutil, hashlib
from datetime import datetime
from urllib.parse import quote
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path('/Users/salvi/MQ/abc-fashion-ai-assistant')
E=ROOT/'docs/test-evidence/d4-final-2026-10-24'
BASE=Path('/private/tmp/d4-final/v1.4/D4 Testing Report Revised.docx')
OUT=ROOT/'output/docx/D4 Testing Report Revised.docx'
PUB=(E/'publication-commit.txt').read_text().strip() if (E/'publication-commit.txt').exists() else 'codex/d4-evidence-links'
REPO='https://github.com/trixielimm/abc-fashion-ai-assistant'
def url(path,tree=False):return REPO+('/tree/' if tree else '/blob/')+PUB+'/'+quote(path,safe='/')
def evidence(file=''):return url('docs/test-evidence/d4-final-2026-10-24'+('/'+file if file else ''),not file)
d=Document(BASE)
P=d.paragraphs
T=d.tables
def setp(i,text):P[i].text=text
def cell(c,text):
    c.text=text
    for p in c.paragraphs:
        for r in p.runs:r.font.size=Pt(9)
def link(p,label,target):
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(target,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');s=OxmlElement('w:rStyle');s.set(qn('w:val'),'Hyperlink');pr.append(s)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'18');pr.append(size);r.append(pr)
    t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)
def result(c,text,file=''):
    cell(c,text+' ');link(c.paragraphs[-1],'[24 Oct evidence]',evidence(file))
def newtable_before(anchor,headers,rows,widths=None):
    table=d.add_table(rows=1,cols=len(headers));table.style=T[2].style
    table.autofit=False
    if widths:
        for col,w in zip(table.columns,widths):col.width=Inches(w)
    for c,s in zip(table.rows[0].cells,headers):cell(c,s)
    for values in rows:
        for c,s in zip(table.add_row().cells,values):cell(c,s)
    # Match the existing report's repeated dark blue header and row spacing.
    for c,source in zip(table.rows[0].cells,T[2].rows[0].cells):
        src=source._tc.tcPr
        if src is not None:
            for tag in ('w:shd','w:tcMar','w:vAlign'):
                x=src.find(qn(tag))
                if x is not None:c._tc.get_or_add_tcPr().append(deepcopy(x))
        for p in c.paragraphs:
            for run in p.runs:run.bold=True;run.font.color.rgb=T[2].rows[0].cells[0].paragraphs[0].runs[0].font.color.rgb
    trPr=table.rows[0]._tr.get_or_add_trPr();rep=OxmlElement('w:tblHeader');trPr.append(rep)
    for row in table.rows:
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
    anchor._p.addprevious(table._tbl)
    return table
def before(i,text,style='Normal'):
    p=P[i].insert_paragraph_before(text,style)
    return p

setp(2,'COMP3850 Deliverable 4   |   Group 20   |   Revised 24 October 2026')
setp(5,'This report specifies and records testing for the eight customer use cases in Figure 7 of the final D4 document (page 62), supported by the UC01–UC07 descriptions on pages 63–72 and the SRS on pages 38–53. It retains stable test identifiers, exact inputs, observed results and evidence links. Browse / Search Products is the eighth customer use case and is identified as BR in this report because the supplied descriptions leave it unnumbered.')
setp(6,'On 24 October, all 48 offline scripts passed; 32 current catalogue/model/HTTP checks and all three isolated AI/database failure/recovery checks passed. The full 38-scenario real-model run produced 108 passes and six provider timeouts; three fresh repetitions of each affected scenario then passed (9/9). Deployed browser testing also found an intermittent occasion conflict, a search recovery defect (fixed locally) and a value-comparison constraint failure. These are retained in DF6–DF8; an all-passing technical readiness decision is not supported. [E14–E16]')
setp(7,"Peyton Bernhardt accepted the product as ready at the 21 October meeting attended by all team members, as confirmed by Salvi. The demonstrated product decision is retained in E13. This report does not substitute that acceptance for technical evidence or Sarah Thomas's final review. The new SRS also contains administrator requirements outside the eight customer use cases; their execution is not established by the customer tests. [E13, E17]")
setp(9,"Report status: revised against all eight customer use cases, with current execution dispositions and disclosed coverage limits. Sarah Thomas is the testing lead and final technical reviewer. The release decision must account for the known intermittent deployed response, unidentified deployed source revision and administrator/browser-platform coverage gaps. No reviewer signature is asserted.")
setp(30,'The customer scope is Browse / Search Products (BR), Shop via Conversation (UC01), Clarify Shopping Preferences (UC02), Obtain Product Recommendations (UC03), Receive Complementary Product Recommendations (UC04), View Product Details (UC05), Refine Shopping Preferences (UC06) and Manage Shopping Cart (UC07). Accessory recommendations, cross-selling and upselling are UC04 branches. Team maintenance/testing and sponsor review/handover in Figure 7 are supporting activities, evaluated through regression, quality checks and the acceptance record.')
setp(32,'Source of truth: final D4 Figure 7, page 62; use case descriptions, pages 63–72; SRS FR-01–FR-40, pages 42–46. The diagram shows eight customer use cases, but the text numbers only UC01–UC07. Table 50 also captions cart management as UC08 while its body says UC07. This report preserves the body identifiers and uses BR for the unnumbered Browse / Search Products case. The matrix below replaces the provisional TO-only mapping. [E17]')
setp(34,'A product can display its one recorded catalogue colour without supporting colour-variant switching. TC-PD-02 and TC-UI-01 now verify accurate fixed-colour presentation for the approved MVP; this is a specification revision, not evidence that the old variant interaction passed. Supplier composition, personal garment fit and per-size stock are not established by aggregate availability or listed sizes. Synthetic product specifications remain labelled. Real checkout, payment and order placement are excluded by the SRS. A complete security audit, sustained load/availability test and native-device browser matrix are not established by these checks.')
setp(37,'Tests follow the revised use case main flows, alternatives and failure conditions. Priorities are explicit size/budget/exclusion constraints, catalogue-supported explanations, correct multi-turn updates and purchase totals, then UI presentation. Each result identifies its execution layer: deterministic fixture, real model with synthetic catalogue, live read-only catalogue, actual HTTP route, deployed browser, or client decision. An assertion at one layer does not establish all the others.')
setp(39,'The final offline run used Node v24.14.0 on macOS; earlier scenario summaries record Node v24.22.1. The checkout uses Next.js 16.3.8, React 19.2.8 and the installed lockfile dependencies. The candidate started at Git HEAD 85d4f1b5e0a22e5dd86445005ebb585055eae7e0 on feat/stylist-ui and includes the retained matching-set/budget, search recovery and harness changes. E14 supplies the source manifest, patch, snapshots and quality-check output; HEAD alone is not the tested candidate.')
setp(40,'Live extraction used the configured OpenAI model; catalogue checks used a read-only public Supabase client and a saved snapshot. Browser actions used the Codex in-app browser at recorded viewports of 1440 × 900, 1280 × 720 and 390 × 844. Its engine version was not exposed by the tool; no Chrome/Safari/Edge or physical iOS/Android certification is claimed. Evidence captures timestamps, exact screenshot dimensions, product IDs and visible state without credentials.')
setp(42,"The named team roles provide responsibility, not an execution signature. Sarah Thomas is the confirmed testing lead. The 24 October scripted and browser tests were launched, observed and documented by Codex at Salvi's request. Sarah's independent review is a separate team activity and has not been represented as completed by this agent.")
setp(46,'D4 is due on 24 October 2026. Dates use Australia/Sydney; raw tool evidence retains UTC timestamps (AEDT = UTC + 11 hours). The schedule below is an execution record with review ownership. Confirm the institutional upload time independently.')
setp(49,'The 21 October client decision remains Accepted. The 24 October technical run, fixes, retests and document/use-case reconciliation are recorded in E14–E17. Earlier milestones remain historical; they are not backdated execution claims.')
setp(52,'Record Git HEAD and all modified source hashes. E14 contains the initial and final local manifests and source patch. Local execution and deployed URL observations have distinct candidate identities; the Vercel source commit was not exposed.')
setp(59,'Recommend technical completion only when prioritised MVP requirements have mapped tests and observed dispositions on the intended candidate; all 48 offline suites pass after relevant changes; and no unresolved critical/high constraint, grounding, arithmetic or failure-disclosure defect remains. Use case mapping is complete for all eight customer cases. Full alternative-flow and SRS compliance require the coverage limitations below to be explicitly reviewed, not silently converted to passes.')
setp(62,'Each record retains case ID, operator (Codex, supervised by the requesting user), date/time, input, expected and actual result, layer, candidate, evidence and defect reference. Independent team review and defect closure are separate. DF identifiers are report references, not invented Jira keys.')
setp(67,'All 45 original test identifiers are retained. The result column now records 24 October observations; historical outcomes remain in the original source E0. PD colour-variant switching is explicitly replaced by fixed-colour display because the revised UC05 requires catalogue-supported details and the implemented product has one colour. No claim of variant-switch support is made.')
setp(68,'Current browser fixture PD1: live catalogue product P1031, Gold Trench Coat, Gold, A$90, sizes XS/S/M/L/XL at https://abcandco.vercel.app/products/P1031. Start with an empty bag. Different sizes of the same product are separate rows; repeated same-size adds increase that row quantity. The recorded desktop/mobile runs use this fixture rather than the earlier mock blouse.')
setp(71,'TC-AI-01–10 use the 24 October live read-only catalogue snapshot: 318 records, including 316 available products. Counts are observations, not fixed acceptance targets. TC-AI-11–14 validate the development-only legacy endpoint; the deployed legacy endpoint correctly returns 404. The production conversation route has separate HTTP cases. [E14]')
setp(74,'TC-PV-01–06 passed within the fresh deterministic validation suite on 24 October. TC-PV-07 separately queried the live Supabase vocabulary. Deterministic and live outcomes are labelled separately. [E14]')
setp(76,'Fresh final execution: 24 October 2026, 18:34:03–18:34:11 AEDT (07:34:03–07:34:11 UTC). All 48 scripts passed after the search recovery change. The source was unchanged during the run. E14 records source hashes and local changes relative to HEAD 85d4f1b. Typecheck, lint and the production webpack build also completed successfully.')
setp(77,'TC-REG-01–46 retain their identities and historical source links. TC-REG-47 and 48 identify the new image-delivery and product-image suites. E14 contains the exact current script/fixture snapshots used on 24 October; use those for current assertions rather than assuming the older E10 source snapshot describes later edits.')
setp(81,'Each scenario starts with its specified fresh session and isolated synthetic catalogue. The harness sends the listed messages in order through the real interpreter plus shortcutPlan, finishConversationTurn and acceptReply. It records each repetition and turn. This is a model/domain integration test; it does not load the browser or query production stock.')
setp(82,'Current assessed scenario outcome: each of the 38 scenarios has three latest passing repetitions. The final full run recorded 108 Pass, 0 Fail and 6 Blocked due to provider connection/timeouts. S3-B, S3-C and S4-A were then each repeated three times, all Pass. Earlier 17 October failures and the first 24 October attempt are retained in the evidence. This is not a claim of one uninterrupted 114/114 run. [E15]')
setp(84,'The suite contains 20 baseline scenarios plus 18 extensions (five budget, six category-default, four collection and three matching-set cases). Initial B5 expected an obsolete optional style question; it was corrected to the implemented grouped-discovery contract. Matching-set wording and budget-only refinement failures were repaired and retested. E15 preserves the initial failures and subsequent outcomes.')
setp(86,'The regression and browser cases cover suggested replies, optional proactive extras, pause/resume, anchor selection, combined budgets, saved-look review, fit guidance, product explanations and value alternatives. These implement UC01–UC06 paths; cart actions are UC07. Current browser results follow below and retain any failed or partially established expectation.')
setp(90,'The following are actual deployed-browser executions on 24 October, with saved accessibility text and screenshots in E16. Pass applies only to the actions and expectations recorded. The fixture prerequisites and model responses are retained; an intermittent failure is not erased by a later passing request.')
setp(92,'The HTTP and isolated dependency failures were executed on 24 October. TC-UI-01 shares the revised fixed-colour scope of TC-PD-02. TC-ACC-01 records the supplied client decision, separately from technical results. There is no assertion of full platform, administrator or deployment-source verification.')
setp(98,'The older Figure 1 records a 22 October navigation answer for historical comparison. The new E16 captures actual desktop/mobile product, cart and AI controls on 24 October. These include quantity arithmetic, persistence, optional-offer controls and saved-look review; none establishes payment or order fulfilment.')
setp(105,"Client decision: accepted as ready on 21 October 2026, as reported by Salvi and transcribed in E13. The final D4 document records the meeting at 6:00 pm (page 55). The URL supplied for current testing is https://abcandco.vercel.app/. The exact source revision demonstrated at that meeting was not supplied. Sarah Thomas owns the final technical review; no signature or review completion is fabricated.")
setp(106,'Handover consists of this editable report, PDF, offline evidence ZIP and the corresponding candidate/evidence snapshots. The new eight-use-case mapping and SRS coverage limits are documented in E17. Product acceptance is retained, while technical limitations remain explicit for the team and assessor.')
setp(108,'E0–E13 retain their fixed historical publication links. E14–E17 link to the new 24 October evidence and use case source in the same GitHub repository. Reviewers require private-repository access; the ZIP contains an offline copy for submission without GitHub authentication. Source references use publication commits; candidate manifests identify the code actually executed.')
setp(110,'Preserve the eight-use-case mapping, revision history and evidence hyperlinks when merging this testing section into D4. The diagram/descriptions are requirement sources; their inclusion does not authorize tests, signatures or acceptance decisions that have not occurred.')

# Replace provisional objective matrix with the eight customer use cases.
uc=[
('BR — Browse / Search Products','Browse catalogue; search by product terms; recover from no results. Diagram-only customer use case.','Figure 7 p.62; web experience','TC-BR-01–03; TC-PD-13; TC-REG search-filters'),
('UC01 — Shop via Conversation','Enter natural-language requirements, obtain recommendations and continue the conversation; disclose failures.','FR-01–03, FR-25–29; pp.63–64','TC-PE-01–08; TC-AI-04; TC-HTTP-01–05; TC-INT-01/02'),
('UC02 — Clarify Shopping Preferences','Ask only for necessary missing/unclear information; accept unknown optional preferences; use latest explicit answers.','FR-03–06, FR-18; pp.64–65','TC-PV-03–06; TC-AIX-01/06; S3-A/B; S6-A/B; guided-discovery suites'),
('UC03 — Obtain Product Recommendations','Retrieve/filter/rank real products; explain matches and unmet soft preferences; preserve hard limits.','FR-07–14, FR-27, FR-30, FR-33; pp.66–67','TC-AI-01–10; TC-INT-03; TC-AIX-07; grounded/validation suites'),
('UC04 — Receive Complementary Product Recommendations','Coordinate multiple roles; provide optional accessories/cross-sells/value alternatives; preserve applicable budget and consent.','FR-21–24; pp.67–68','TC-AIX-02–05/08; B1–B5; S4-A/B; S5-A–F; cross-sell/upsell/budget suites'),
('UC05 — View Product Details','Open selected product; display supported image, price, colour, size and availability; handle missing details/images.','FR-11, FR-13, FR-31, FR-33; pp.69–70','TC-PD-01–03/06/07/12; TC-UI-01; TC-AIX-07; TC-REG-47/48'),
('UC06 — Refine Shopping Preferences','Change one stated preference, retain others, use newest explicit value, reject/replace recommendations or clarify ambiguity.','FR-15–20, FR-28; pp.70–71','S1-A–D; S2-A–D; S3-C; CAT/COL/SET; cheaper/scoped-state suites'),
('UC07 — Manage Shopping Cart','Select size, add available item, inspect cart, change quantity/remove, retain cart and display empty state.','FR-32; pp.71–72 (Table 50 caption says UC08)','TC-PD-04/05/08–14; TC-UI-02; cart and failure boundary checks'),
]
for row in list(T[2].rows)[1:]:T[2]._tbl.remove(row._tr)
for c,s in zip(T[2].rows[0].cells,['Use case','Behaviour under test','Requirement / source','Test coverage']):cell(c,s)
for row in uc:
    for c,s in zip(T[2].add_row().cells,row):cell(c,s)

# Revision history and execution schedule.
for row in T[1].rows[1:]:
    if row.cells[0].text=='Not Run':
        cell(row.cells[0],'Coverage gap');cell(row.cells[1],'A required behaviour has no executed evidence here. It is disclosed explicitly and is not counted as a pass.')
cell(T[4].rows[0].cells[1],'Use and evidence boundary')
cell(T[4].rows[1].cells[1],'Read-only public Supabase client retrieved the live snapshot. No shared catalogue record was changed; no dashboard CRUD execution is claimed.')
cell(T[4].rows[2].cells[1],'Browser screenshots, accessibility state and captured console logs accompany observed actions. The embedded browser engine version was not exposed.')
cell(T[4].rows[3].cells[1],'Confirmed submission URL https://abcandco.vercel.app/: HTTP and browser behaviour observed. Source commit/deployment equality with local fixes is unverified.')
cell(T[5].rows[-1].cells[-1],'Codex launched and observed the 24 October scripts and browser actions at Salvi’s request. Sarah remains the testing lead and independent technical reviewer.')
cell(T[7].rows[3].cells[-1],'Uses configured model credentials and synthetic catalogue; three repetitions per scenario were executed. Save every outcome. It does not test HTTP or rendering.')
cell(T[7].rows[5].cells[1],'npm run typecheck; npm run lint; npm run build -- --webpack')
cell(T[7].rows[5].cells[-1],'Separate checks, all completed with exit code0 on 24 October after the local search recovery fix. See E14 logs.')
cell(T[8].rows[-1].cells[0],'Single colour and per-size stock limits')
cell(T[8].rows[-1].cells[1],'Test the actual recorded colour and listed sizes. The revised fixed-colour case does not establish variant switching or stock for a particular size.')
for row in T[17].rows[1:]:
    s=row.cells[0].text
    s=re.sub(r'TO[0-9]+(?:, TO[0-9]+)*','',s).strip()
    if 'Guided discovery' in s:s+=' — UC01/UC02'
    elif 'Proactive' in s:s+=' — UC04'
    elif 'Suggested outfit' in s:s+=' — UC04/UC06'
    elif 'Explanations' in s:s+=' — UC02/UC03/UC05'
    elif 'Cross-sell' in s:s+=' — UC04'
    else:s+=' — UC06/UC04'
    cell(row.cells[0],s)
r=T[0].add_row();
for c,s in zip(r.cells,['V1.5 — 24 October 2026','Codex at Salvi’s request; Sarah Thomas remains testing lead','Aligned with eight customer use cases in final D4 pp.62–72 and SRS; executed current tests, retained failures/retests and replaced provisional statuses with observed dispositions.']):cell(c,s)
for row in T[6].rows[1:]:
    text=' '.join(c.text for c in row.cells)
    if '22 Oct' in text:cell(row.cells[-1],'Completed historical baseline; superseded by the 24 October run.')
    elif '23 Oct' in text:cell(row.cells[-1],'Reconciled on 24 October using the newly supplied final D4 document; original planned date retained.')
    elif '24 Oct' in text:cell(row.cells[-1],'Execution and report revision completed as recorded; independent technical review remains Sarah’s responsibility.')

# Actual original-case results; no conversion of old unexecuted variant tests.
pd={
1:'Pass — P1031 details loaded; Gold, A$90 and XS/S/M/L/XL matched the live fixture.',
2:'Pass — revised fixed-colour test: Gold label/swatch matched P1031. One catalogue colour exists; colour switching is outside this case’s revised MVP scope.',
3:'Pass — selected S and then M; selection state changed correctly.',
4:'Pass — selected S, added P1031; bag showed Gold/S at A$90.',
5:'Pass — Add without size displayed “Please select a size”; bag count remained zero.',
6:'Pass — Details expanded to recorded construction/fit/style/occasion information and collapsed on the next activation.',
7:'Pass — 390 × 844 product view retained readable details, size controls and Add action without blocking overlap.',
8:'Pass — S quantity1 and M quantity2 were separate rows; total units3 and subtotal A$270.',
9:'Pass — repeated M add increased its row quantity to2; S1+M2 total A$270.',
10:'Pass — removing S left M2, bag count2 and A$180; final mobile removal also produced the empty-cart state.',
11:'Pass — M quantity2→3→2 recalculated A$180→270→180. Reload retained product, M and quantity2; mobile persistence also verified.',
12:'Pass — P1031 link from bag opened the matching Gold Trench Coat product page at A$90.',
13:'Pass — Continue Browsing opened the collection/homepage and retained the bag.',
14:'Pass — cart styling-advice control opened AI Stylist in context.',
}
for n,row in enumerate(T[9].rows[1:],1):
    cell(row.cells[1],row.cells[1].text.replace('/products/silk-blend-blouse','/products/P1031').replace('Ivory','Gold').replace('A$258','A$180').replace('A$129','A$90'))
    result(row.cells[3],pd[n])
cell(T[9].rows[2].cells[0],'TC-PD-02 — Recorded colour display (UC05; scope revised)')
cell(T[9].rows[2].cells[1],'Open P1031; compare the Gold catalogue colour with the visible swatch/label and saved cart row. Do not require nonexistent variants.')
cell(T[9].rows[2].cells[2],'The one recorded colour is displayed accurately and consistently; the UI makes no unsupported variant promise.')
cell(T[9].rows[4].cells[1],'On P1031 select S and Add to bag; inspect Gold/S and A$90 in the cart.')
cell(T[9].rows[14].cells[2],'The contextual AI Stylist interface opens for further shopping assistance.')

http=json.loads((E/'http-model-catalogue-latest.json').read_text())
byid={c['id']:c for c in http['checks']}
for n,row in enumerate(T[10].rows[1:],1):
    key=f'TC-AI-{n:02}'
    if n<=10:
        c=byid[key];actual=c.get('actual')
        text='Pass — fresh live catalogue/model/domain execution met the stated assertion.'
        if n==1:text='Pass — read 318 current catalogue records (316 available); snapshot retained.'
        if n==2:text=f'Pass — {actual["count"]} available/in-stock records met size10 and maximum A$100.'
        if n==10:text='Pass — no eligible products at maximum A$1; hard-filter result was empty.'
        result(row.cells[3],text,'http-model-catalogue-latest.json')
    elif n<=14:result(row.cells[3],'Pass — current development HTTP validation run returned the expected 400/415 and controlled error contract. Production legacy endpoint was separately confirmed 404.','legacy-api-validation.txt')
    else:
        result(row.cells[3],'Pass — actual POST route in an isolated HTTP process returned '+('502 CONVERSATION_FAILED' if n==15 else '503 CATALOGUE_UNAVAILABLE')+'; safe message, no fabricated products or shared-service disruption. See TC-INT-01/02.','controlled-failures.json')
for n,row in enumerate(T[11].rows[1:],1):result(row.cells[3],'Pass — 24 October fresh '+('real-model extraction matched all specified fields and preserved unspecified values.' if n<=6 else 'blank-input/HTTP validation returned the specified controlled rejection.'),'http-model-catalogue-latest.json')
for n,row in enumerate(T[12].rows[1:],1):result(row.cells[3],'Pass — '+('fresh deterministic validation assertions passed in the 48-suite run.' if n<=6 else 'fresh live vocabulary recognised Brunch and retained a non-blocking unsupported Casual style issue.'),'offline-final.txt' if n<=6 else 'http-model-catalogue-latest.json')

# Current offline suite identity/source links.
oldfiles=[]
for row in T[13].rows[1:]:
    name=row.cells[1].text.strip();oldfiles.append(name)
    result(row.cells[3],'Pass — 24 Oct offline run.','offline-final.txt')
    cell(row.cells[1],'');link(row.cells[1].paragraphs[0],name,evidence('source/scripts/'+name))
for num,name in [(47,'test-image-delivery.ts'),(48,'test-product-images.ts')]:
    row=T[13].add_row();cell(row.cells[0],f'TC-REG-{num:02}');cell(row.cells[1],'');link(row.cells[1].paragraphs[0],name,evidence('source/scripts/'+name));cell(row.cells[2],'Current source snapshot: image delivery/fallback and catalogue product image assertions.');result(row.cells[3],'Pass — 24 Oct offline run.','offline-final.txt')
for table in [T[15],T[16]]:
    for row in table.rows[1:]:
        cid=re.match(r'^[A-Z0-9]+(?:-[A-Z0-9]+)?',row.cells[0].text).group()
        result(row.cells[3],'Pass — three latest real-model/fixture repetitions. '+('Provider timeouts in the full run were followed by three passing fresh repetitions; original failures retained.' if cid in ['S3-B','S3-C','S4-A'] else 'See E15 for all attempts and turn-level assertions.'))

# UI outcomes supplied by the actual browser record. Updated after testing.
ui=json.loads((E/'report-ui-outcomes.json').read_text()) if (E/'report-ui-outcomes.json').exists() else {}
for n,row in enumerate(T[18].rows[1:],1):
    result(row.cells[3],ui.get(f'TC-AIX-{n:02}','Execution result requires reconciliation with saved browser evidence.'))

for row in T[19].rows[1:]:
    cid=row.cells[0].text.split()[0]
    if cid.startswith('TC-HTTP-'):result(row.cells[3],'Pass — local and deployed endpoints returned the specified controlled HTTP status/body. '+('Valid request at exactly64000 bytes returned200; 64001 returned413.' if '05' in cid else ''),'http-model-catalogue-latest.json')
    elif cid.startswith('TC-INT-01'):result(row.cells[3],'Pass — injected model failure returned502 CONVERSATION_FAILED, safe message and unchanged request state.','controlled-failures.json')
    elif cid.startswith('TC-INT-02'):result(row.cells[3],'Pass — injected catalogue failure returned503 CATALOGUE_UNAVAILABLE; restored service returned200 with catalogue-backed products.','controlled-failures.json')
    elif cid.startswith('TC-UI-01'):
        cell(row.cells[1],'Revised scope: compare P1031 Gold colour label/swatch, product page and cart row; verify only supported catalogue information. The product has no alternate colour variant.')
        cell(row.cells[2],'Recorded single colour is consistent across product and cart; no variant-switch functionality is claimed.')
        result(row.cells[3],'Pass — fixed-colour display verified. Original multicolour interaction is superseded by this explicit scope revision.')
    elif cid.startswith('TC-UI-02'):result(row.cells[3],'Pass — desktop/mobile add, separate-size rows, repeat add, increase/decrease, removal, navigation and reload persistence completed with P1031 A$90. Empty cart restored.')
    elif cid.startswith('TC-INT-03'):result(row.cells[3],'Pass — deployed product IDs, names, prices and colours matched the saved current catalogue snapshot. Listed sizes do not assert per-size stock.','http-model-catalogue-latest.json')
    elif cid.startswith('TC-DEP-01'):result(row.cells[3],ui.get('TC-DEP-01','Partial — deployed root HTTP200 and live conversation response observed. Deployed source commit could not be identified; this is a build-traceability coverage limit.'))

# Replace execution summary, retain the old evidence register as historical.
for row in list(T[20].rows)[1:]:T[20]._tbl.remove(row._tr)
for vals in [
('Offline candidate','48/48 suites Pass on24Oct; typecheck, lint and production webpack build Pass. Local changes and hashes are retained in E14.'),
('Catalogue/model/HTTP','32/32 latest check IDs Pass. Initial request-context setup failures retained in attempt-1; corrected standalone read-only catalogue setup then passed.'),
('Dependency failures','3/3 isolated HTTP checks Pass: provider502, catalogue503 and recovery200. Shared credentials/services were unchanged.'),
('Real-model customer scenarios','Full final run108Pass/0Fail/6provider-timeoutBlocked; affected three scenarios repeated3times,9/9Pass. All38scenarios have three latest passing repetitions; all attempts retained.'),
('Deployed browser','Product/cart and recorded guidance/offer controls passed. AIX04 has an intermittent false conflict, BR03 failed (fixed locally) and AIX08 value comparison failed to preserve the saved budget. See DF6–DF8.'),
('Client decision','Accepted21Oct2026, reported by Salvi; final D4 records6pm. Sarah remains responsible for independent technical review.'),
('Coverage boundary','Eight customer use cases mapped. Full SRS certification is not claimed: administrator FR34–40, native-browser matrix, sustained NFR thresholds and several mutation/error alternatives lack executed evidence.'),
]:
    for c,s in zip(T[20].add_row().cells,vals):cell(c,s)
defects=[
('DF1','Historical dependency/setup issue','Resolved for current offline execution:48/48Pass. First live-catalogue harness used a request-bound helper outside HTTP; corrected setup passed. First attempt retained.'),
('DF2','Historical hard-constraint and refinement failures','Current38scenario suite has three latest passing repetitions each; six provider timeouts were rerun successfully. Does not establish browser reliability for every free-text phrasing.'),
('DF3','Matching-set wording and budget-only refinement','Local targeted fix verified by named regression assertions,12/12targeted model runs and the complete scenario run/recovery. Source patch retained; deployed source equality is unverified.'),
('DF4','Service failures and old multicolour test','AI/database/recovery checks Pass. Multicolour requirement explicitly replaced with recorded single-colour display, aligned to UC05; no variant feature was added.'),
('DF5','Provisional SRS/use-case mapping','Reconciled against final D4Figure7pp.62–72 and FR01–40. Eightcustomercases mapped; administrator and other unproved alternatives are disclosed in the coverage limits.'),
('DF6','Intermittent false occasion conflict on deployed combined request','Retained failure: dress+bag/beach-wedding/A$200 request reported conflicting saved occasion on a new chat. Fresh local and deployed HTTP probes later returned valid A$155–195 combinations. Cause/consistent deployed repair not established; treat as known intermittent limitation.'),
('DF7','Browse all products did not clear deployed search state','Fail on deployed BR03: URL changed but no-results state remained. SearchExplorer now resets query and reloads through the controller. Local browser retest displayed316 products; quality checks passed. This local fix is not asserted deployed.'),
('DF8','Value alternatives lost saved maximum budget','Fail on deployed AIX08: exact value-comparison request was interpreted as occasion text and cards atA$200/A$240 appeared after the savedA$150 maximum. Budget/intent regression requires a fix and deployed confirmation; current technical acceptance criterion is not met.'),
]
for row in list(T[21].rows)[1:]:T[21]._tbl.remove(row._tr)
for vals in defects:
    for c,s in zip(T[21].add_row().cells,vals):cell(c,s)

# Explicit main/alternative coverage, including gaps rather than false passes.
before(35,'Use case flow coverage and limits','Heading 2')
before(35,'The eight-case matrix establishes traceability. The following flow review distinguishes executed behaviours from requirements that the current evidence does not establish. A coverage gap is not an unexecuted test disguised as a pass. Detailed requirement IDs remain those in the supplied D4 source.')
flowrows=[
('BR','Search/browse, exact product-name match, no-results message and recovery; search filter logic has offline assertions.','TC-BR-01–03 and search-filters suite. Browser filter/pagination breadth is limited to the actions recorded.'),
('UC01','Main conversation and UC03 handoff; clarification and controlled technical failures.','Pass at the recorded model/domain, browser and HTTP layers; not a guarantee for arbitrary language.'),
('UC02','Missing preferences, unknown optional size, known/conflicting constraints and latest explicit values.','Guided-discovery, validation, scope and repeated scenarios Pass. Human judgement of every unclear paraphrase is not established.'),
('UC03','Live filtering/ranking, grounded details/explanations, soft alternatives, hard no-match, unknown-size browsing.','TC-AI/PV/INT and AIX Pass for stated inputs. No complete catalogue-wide recommendation relevance audit is claimed.'),
('UC04','Optional extras, anchor switch, pause/resume, combined roles/budgets, saved look and value alternatives.','See AIX outcomes. One intermittent deployed occasion-conflict failure retained as DF6; later successes do not erase it.'),
('UC05','Main details/navigation; incomplete attributes and missing-image fallback have deterministic coverage.','Fixed-colour browser case Pass. Live mutation of product availability and live catalogue retrieval-error UI alternatives were not injected.'),
('UC06','Cheaper/colour/category/collection/budget changes, preserved context, rejection and latest-value scope.','Repeated scenario and deterministic suites Pass. Explicit global constraints are retained until explicitly changed.'),
('UC07','Add, required-size prompt, view, separate variants bysize, quantity, remove, persistence and empty state.','Desktop/mobile browser Pass. Availability changing between display and add, and client-cart storage failure alternatives are coverage gaps; no write was made to shared production stock.'),
('Supporting roles / SRS','Sponsor acceptance; team regression/build and requirements review.','Client acceptance recorded. AdminFR34–40 are additional SRS requirements and have no demonstrated CRUD/auth execution here; the customer diagram does not establish their implementation.'),
('Non-functional requirements','Responsive views, controlled errors and relevant quality checks.','Physical-device Safari/Chrome/Edge, sustained response-time/99%availability, full security and load thresholds are not certified by these runs.'),
]
newtable_before(P[35],['Use case / scope','Flows addressed','Evidence disposition / limit'],flowrows,[0.8,2.7,2.9])

before(75,'Browse / Search Products — eighth customer use case','Heading 2')
before(75,'BR is a report alias for the unnumbered customer function in Figure7. These cases are separate from conversational discovery and retain their own browser evidence.')
brrows=[
('TC-BR-01 — product search','Open Shop; query “Gold Trench Coat”.','Matching catalogue cards and links; accurate names/prices; bag unchanged.','Pass — P1031A$90 and P1093A$215 returned on mobile.'),
('TC-BR-02 — no results','Search “zzzd4nomatch2026”.','Zero products, understandable no-results message and recovery action.','Pass — 0products; “No products found”; broader-term guidance and Browse all products link.'),
('TC-BR-03 — recovery','Activate Browse all products after no results.','Return to the unfiltered collection with selectable product cards.','Fail deployed: query/results remained. Pass local retest after fix: cleared query,316 products and restored cards. DF7; deployment not established.'),
]
bt=newtable_before(P[75],['Case / use case','Input / action','Expected output','Actual output'],brrows,[1.0,1.65,1.8,1.95])
for row in bt.rows[1:]:link(row.cells[3].paragraphs[0],' [E16]',evidence())

for vals in [
('E14','24October candidate, offline/HTTP/catalogue/model checks, controlled failures, quality checks and current source snapshots',evidence()),
('E15','All 24October customer-scenario attempts and later timeout recovery; 38cases with three latest passing repetitions',url('docs/test-evidence',True)),
('E16','24October browser-results.jsonl plus per-case screenshots and accessibility captures',evidence('browser-results.jsonl')),
('E17','Final supplied D4 document: Figure7p.62, UCdescriptionspp.63–72, SRSpp.38–53; eight-case alignment and explicit coverage limits',evidence('use-case-alignment.md')),
]:
    row=T[23].add_row();cell(row.cells[0],vals[0]);cell(row.cells[1],'');link(row.cells[1].paragraphs[0],vals[1],vals[2])

# Update residual version/date tokens in footer/header only.
for sec in d.sections:
    for part in [sec.header,sec.footer]:
        for p in part.paragraphs:
            for run in p.runs:
                run.text=run.text.replace('v1.4','v1.5').replace('V1.4','V1.5').replace('22 October 2026','24 October 2026')
OUT.parent.mkdir(parents=True,exist_ok=True)
d.save(OUT)
print(OUT)
